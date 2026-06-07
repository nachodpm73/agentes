#!/usr/bin/env python3
"""
Extrae de Mailchimp los datos de seguimiento de uno o varios publicos para un
periodo (por defecto marzo-abril 2026) y vuelca un EXCEL por publico con todas
las tablas que necesita el informe ejecutivo:

  - Resumen          (tasa de lectura, lecturas unicas, clics, bajas, rebotes, ...)
  - Campanas         (campana a campana: asunto, fecha, env., apert., %lect, clics, %clics, bajas)
  - Mensual          (comparativa por mes)
  - Top5             (top 5 campanas por tasa de lectura)
  - ListaCaliente    (segmentacion de contactos por nº de aperturas: >=20 / 10-19 / <10)
  - TopEmpresas      (top 10 dominios por aperturas, excluyendo genericos)
  - SaludBase        (bajas, rebotes duros/suaves, dominios sin aperturas, reduccion neta)

Uso:
    set MAILCHIMP_API_KEY=xxxxxxxx-usXX
    py extraer_datos_publicos.py --desde 2026-03-01 --hasta 2026-04-30

Por defecto procesa: OPERAND, OPERBCN, OPERMAD, OPERVAL, OPERBIO, OPERGAL, OPERCAN
"""

import argparse
import base64
import datetime as dt
import os
import sys
import time

import requests
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

PUBLICOS = ["OPERAND", "OPERBCN", "OPERMAD", "OPERVAL", "OPERBIO", "OPERGAL", "OPERCAN"]

# Dominios genericos y propios que se excluyen del ranking de empresas
GENERICOS = {
    "gmail.com", "hotmail.com", "hotmail.es", "outlook.com", "outlook.es",
    "yahoo.com", "yahoo.es", "live.com", "icloud.com", "me.com", "aol.com",
    "msn.com", "terra.es", "telefonica.net", "operinter.com",
}


class Mailchimp:
    def __init__(self, api_key):
        if "-" not in api_key:
            sys.exit("ERROR: la API key de Mailchimp no tiene el sufijo -usXX.")
        self.dc = api_key.split("-")[-1]
        self.base = f"https://{self.dc}.api.mailchimp.com/3.0"
        self.s = requests.Session()
        token = base64.b64encode(f"anystring:{api_key}".encode()).decode()
        self.s.headers.update({"Authorization": f"Basic {token}"})

    def get(self, path, params=None):
        for attempt in range(6):
            r = self.s.get(self.base + path, params=params, timeout=90)
            if r.status_code == 429:
                time.sleep(2 ** attempt)
                continue
            if not r.ok:
                sys.exit(f"ERROR Mailchimp {r.status_code} en {path}: {r.text[:300]}")
            return r.json()
        sys.exit(f"ERROR: rate limit persistente en {path}")

    def find_list(self, name):
        data = self.get("/lists", {"count": 1000, "fields": "lists.id,lists.name"})
        nl = name.lower()
        for l in data.get("lists", []):
            if l["name"].lower() == nl:
                return l
        for l in data.get("lists", []):
            if nl in l["name"].lower():
                return l
        return None

    def campaigns_in_range(self, list_id, desde, hasta):
        """Campanas enviadas entre desde y hasta (inclusive)."""
        out, offset = [], 0
        hasta_excl = (dt.date.fromisoformat(hasta) + dt.timedelta(days=1)).isoformat()
        while True:
            data = self.get("/campaigns", {
                "list_id": list_id, "status": "sent", "count": 500, "offset": offset,
                "since_send_time": f"{desde}T00:00:00+00:00",
                "before_send_time": f"{hasta_excl}T00:00:00+00:00",
                "sort_field": "send_time", "sort_dir": "ASC",
                "fields": "campaigns.id,campaigns.send_time,"
                          "campaigns.settings.subject_line,campaigns.settings.title,total_items",
            })
            batch = data.get("campaigns", [])
            out.extend(batch)
            offset += len(batch)
            if not batch or offset >= data.get("total_items", 0):
                break
        return out

    def report(self, campaign_id):
        return self.get(f"/reports/{campaign_id}", {
            "fields": "emails_sent,unsubscribed,bounces,opens,clicks,abuse_reports",
        })

    def email_activity(self, campaign_id):
        """Devuelve dict {email: nº aperturas} para una campana."""
        opens, offset = {}, 0
        while True:
            data = self.get(f"/reports/{campaign_id}/email-activity", {
                "count": 1000, "offset": offset,
                "fields": "emails.email_address,emails.activity.action,total_items",
            })
            emails = data.get("emails", [])
            for e in emails:
                addr = e.get("email_address", "").lower()
                n = sum(1 for a in e.get("activity", []) if a.get("action") == "open")
                if n:
                    opens[addr] = opens.get(addr, 0) + n
            offset += len(emails)
            if not emails or offset >= data.get("total_items", 0):
                break
        return opens

    def subscribed_domains(self, list_id):
        """Conjunto de dominios de los miembros suscritos."""
        doms, offset = set(), 0
        while True:
            data = self.get(f"/lists/{list_id}/members", {
                "status": "subscribed", "count": 1000, "offset": offset,
                "fields": "members.email_address,total_items",
            })
            members = data.get("members", [])
            for m in members:
                addr = m.get("email_address", "")
                if "@" in addr:
                    doms.add(addr.split("@")[1].lower())
            offset += len(members)
            if not members or offset >= data.get("total_items", 0):
                break
        return doms


def pct(n, d):
    return round(100.0 * n / d, 2) if d else 0.0


def procesar_publico(mc, nombre, desde, hasta):
    lst = mc.find_list(nombre)
    if not lst:
        print(f"  [{nombre}] AVISO: audiencia no encontrada, se omite.")
        return None
    list_id = lst["id"]
    camps = mc.campaigns_in_range(list_id, desde, hasta)
    print(f"  [{nombre}] {len(camps)} campanas en {desde}..{hasta}")
    if not camps:
        return {"nombre": nombre, "list_id": list_id, "campanas": []}

    filas = []
    opens_acum = {}  # email -> aperturas en el bimestre
    tot = {"env": 0, "apert": 0, "clics": 0, "bajas": 0, "hard": 0, "soft": 0}
    for c in camps:
        rep = mc.report(c["id"])
        env = rep.get("emails_sent", 0)
        apert = rep.get("opens", {}).get("unique_opens", 0)
        clics = rep.get("clicks", {}).get("unique_subscriber_clicks",
                rep.get("clicks", {}).get("unique_clicks", 0))
        bajas = rep.get("unsubscribed", 0)
        hard = rep.get("bounces", {}).get("hard_bounces", 0)
        soft = rep.get("bounces", {}).get("soft_bounces", 0)
        fecha = (c.get("send_time", "")[:10])
        try:
            fecha = dt.datetime.fromisoformat(c["send_time"].replace("Z", "+00:00")).strftime("%d/%m/%Y")
        except Exception:
            pass
        filas.append({
            "asunto": c.get("settings", {}).get("subject_line", ""),
            "fecha": fecha, "send_time": c.get("send_time", ""),
            "env": env, "apert": apert, "pct_lect": pct(apert, env),
            "clics": clics, "pct_clics": pct(clics, env), "bajas": bajas,
        })
        for k, v in (("env", env), ("apert", apert), ("clics", clics),
                     ("bajas", bajas), ("hard", hard), ("soft", soft)):
            tot[k] += v
        for addr, n in mc.email_activity(c["id"]).items():
            opens_acum[addr] = opens_acum.get(addr, 0) + n

    # Segmentacion lista caliente
    con_apertura = {a: n for a, n in opens_acum.items() if n > 0}
    muy = sum(1 for n in con_apertura.values() if n >= 20)
    comp = sum(1 for n in con_apertura.values() if 10 <= n < 20)
    ocas = sum(1 for n in con_apertura.values() if n < 10)

    # Top empresas por dominio
    dom_opens = {}
    for addr, n in con_apertura.items():
        if "@" not in addr:
            continue
        d = addr.split("@")[1].lower()
        if d in GENERICOS:
            continue
        dom_opens[d] = dom_opens.get(d, 0) + n
    top_emp = sorted(dom_opens.items(), key=lambda x: x[1], reverse=True)[:10]

    # Salud base: dominios sin aperturas
    doms_base = mc.subscribed_domains(list_id)
    doms_con_apertura = {a.split("@")[1].lower() for a in con_apertura if "@" in a}
    doms_sin = len(doms_base - doms_con_apertura)

    base_ini = filas[0]["env"] if filas else 0
    base_fin = filas[-1]["env"] if filas else 0

    # Mensual
    meses = {}
    for f in filas:
        try:
            m = dt.datetime.fromisoformat(f["send_time"].replace("Z", "+00:00")).strftime("%Y-%m")
        except Exception:
            m = f["fecha"][3:10]
        d = meses.setdefault(m, {"camp": 0, "env": 0, "apert": 0})
        d["camp"] += 1; d["env"] += f["env"]; d["apert"] += f["apert"]

    return {
        "nombre": nombre, "list_id": list_id, "campanas": filas, "tot": tot,
        "pct_lect_media": pct(tot["apert"], tot["env"]),
        "pct_clics_media": pct(tot["clics"], tot["env"]),
        "pct_bajas": pct(tot["bajas"], tot["env"]),
        "con_apertura": len(con_apertura), "base_fin": base_fin, "base_ini": base_ini,
        "muy": muy, "comp": comp, "ocas": ocas,
        "top_emp": top_emp, "num_camp": len(filas),
        "doms_sin": doms_sin, "doms_base": len(doms_base),
        "meses": dict(sorted(meses.items())),
    }


# --------------------------------------------------------------------------- #
# Excel
# --------------------------------------------------------------------------- #
H_FILL = PatternFill("solid", fgColor="1F3864")
H_FONT = Font(bold=True, color="FFFFFF")


def _head(ws, row, headers):
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=j, value=h)
        c.fill = H_FILL; c.font = H_FONT; c.alignment = Alignment(horizontal="center")


def escribir_excel(d, periodo, outdir):
    wb = Workbook()

    ws = wb.active; ws.title = "Resumen"
    ws["A1"] = f"Informe de seguimiento - {d['nombre']}"; ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Periodo: {periodo}"
    res = [
        ("Campanas", d["num_camp"]),
        ("Envios totales", d["tot"]["env"]),
        ("Lecturas unicas", d["tot"]["apert"]),
        ("Tasa de lectura media", f"{d['pct_lect_media']}%"),
        ("Clics unicos", d["tot"]["clics"]),
        ("Tasa de clics", f"{d['pct_clics_media']}%"),
        ("Bajas", d["tot"]["bajas"]),
        ("Tasa de bajas", f"{d['pct_bajas']}%"),
        ("Rebotes duros (invalidos)", d["tot"]["hard"]),
        ("Rebotes suaves", d["tot"]["soft"]),
        ("Contactos con >=1 apertura", d["con_apertura"]),
        ("  - Muy comprometidos (>=20)", d["muy"]),
        ("  - Comprometidos (10-19)", d["comp"]),
        ("  - Ocasionales (<10)", d["ocas"]),
        ("Base al inicio / cierre", f"{d['base_ini']} / {d['base_fin']}"),
        ("Reduccion neta de la base", d["base_fin"] - d["base_ini"]),
        ("Dominios sin ninguna apertura", d["doms_sin"]),
    ]
    for i, (k, v) in enumerate(res, 4):
        ws.cell(row=i, column=1, value=k); ws.cell(row=i, column=2, value=v)
    ws.column_dimensions["A"].width = 34; ws.column_dimensions["B"].width = 22

    ws = wb.create_sheet("Campanas")
    _head(ws, 1, ["Asunto", "Fecha", "Env.", "Apert.", "% Lect.", "Clics", "% Clics", "Bajas"])
    for i, f in enumerate(d["campanas"], 2):
        ws.cell(row=i, column=1, value=f["asunto"])
        ws.cell(row=i, column=2, value=f["fecha"])
        ws.cell(row=i, column=3, value=f["env"])
        ws.cell(row=i, column=4, value=f["apert"])
        ws.cell(row=i, column=5, value=f["pct_lect"] / 100).number_format = "0.0%"
        ws.cell(row=i, column=6, value=f["clics"])
        ws.cell(row=i, column=7, value=f["pct_clics"] / 100).number_format = "0.00%"
        ws.cell(row=i, column=8, value=f["bajas"])
    r = len(d["campanas"]) + 2
    ws.cell(row=r, column=1, value="TOTALES / MEDIA").font = Font(bold=True)
    ws.cell(row=r, column=3, value=d["tot"]["env"]).font = Font(bold=True)
    ws.cell(row=r, column=4, value=d["tot"]["apert"]).font = Font(bold=True)
    ws.cell(row=r, column=5, value=d["pct_lect_media"] / 100).number_format = "0.0%"
    ws.cell(row=r, column=6, value=d["tot"]["clics"]).font = Font(bold=True)
    ws.cell(row=r, column=7, value=d["pct_clics_media"] / 100).number_format = "0.00%"
    ws.cell(row=r, column=8, value=d["tot"]["bajas"]).font = Font(bold=True)
    ws.column_dimensions["A"].width = 46
    for col in "BCDEFGH":
        ws.column_dimensions[col].width = 10

    ws = wb.create_sheet("Mensual")
    _head(ws, 1, ["Mes", "Campanas", "Envios", "Lecturas", "% Lectura"])
    for i, (m, v) in enumerate(d["meses"].items(), 2):
        ws.cell(row=i, column=1, value=m)
        ws.cell(row=i, column=2, value=v["camp"])
        ws.cell(row=i, column=3, value=v["env"])
        ws.cell(row=i, column=4, value=v["apert"])
        ws.cell(row=i, column=5, value=pct(v["apert"], v["env"]) / 100).number_format = "0.0%"
    for col in "ABCDE":
        ws.column_dimensions[col].width = 14

    ws = wb.create_sheet("Top5")
    _head(ws, 1, ["Pos.", "Campana", "Fecha", "Envios", "% Lectura", "Clics"])
    top5 = sorted(d["campanas"], key=lambda f: f["pct_lect"], reverse=True)[:5]
    for i, f in enumerate(top5, 2):
        ws.cell(row=i, column=1, value=i - 1)
        ws.cell(row=i, column=2, value=f["asunto"])
        ws.cell(row=i, column=3, value=f["fecha"])
        ws.cell(row=i, column=4, value=f["env"])
        ws.cell(row=i, column=5, value=f["pct_lect"] / 100).number_format = "0.0%"
        ws.cell(row=i, column=6, value=f["clics"])
    ws.column_dimensions["B"].width = 46

    ws = wb.create_sheet("ListaCaliente")
    _head(ws, 1, ["Nivel", "Contactos", "% del total", "Aperturas", "Perfil"])
    tot_ap = d["con_apertura"] or 1
    niveles = [
        ("Muy comprometidos", d["muy"], ">= 20", "Lista Caliente"),
        ("Comprometidos", d["comp"], "10 - 19", "Seguimiento"),
        ("Ocasionales", d["ocas"], "< 10", "Observacion"),
    ]
    for i, (n, c, ap, pf) in enumerate(niveles, 2):
        ws.cell(row=i, column=1, value=n)
        ws.cell(row=i, column=2, value=c)
        ws.cell(row=i, column=3, value=pct(c, tot_ap) / 100).number_format = "0.0%"
        ws.cell(row=i, column=4, value=ap)
        ws.cell(row=i, column=5, value=pf)
    ws.cell(row=5, column=1, value="TOTAL CON APERTURAS").font = Font(bold=True)
    ws.cell(row=5, column=2, value=d["con_apertura"]).font = Font(bold=True)
    ws.column_dimensions["A"].width = 22; ws.column_dimensions["E"].width = 16

    ws = wb.create_sheet("TopEmpresas")
    _head(ws, 1, ["Pos.", "Empresa (dominio)", "Total Aperturas", "Media por Campana"])
    for i, (dom, ap) in enumerate(d["top_emp"], 2):
        ws.cell(row=i, column=1, value=i - 1)
        ws.cell(row=i, column=2, value=dom)
        ws.cell(row=i, column=3, value=ap)
        ws.cell(row=i, column=4, value=round(ap / (d["num_camp"] or 1), 1))
    ws.column_dimensions["B"].width = 34; ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 18

    ws = wb.create_sheet("SaludBase")
    _head(ws, 1, ["Indicador", "Valor"])
    salud = [
        ("Bajas voluntarias", f"{d['tot']['bajas']} ({d['pct_bajas']}%)"),
        ("Rebotes duros (invalidos)", d["tot"]["hard"]),
        ("Rebotes suaves (temporales)", d["tot"]["soft"]),
        ("Dominios sin ninguna apertura", d["doms_sin"]),
        ("Dominios en base", d["doms_base"]),
        ("Reduccion neta de la base", d["base_fin"] - d["base_ini"]),
    ]
    for i, (k, v) in enumerate(salud, 2):
        ws.cell(row=i, column=1, value=k); ws.cell(row=i, column=2, value=v)
    ws.column_dimensions["A"].width = 32; ws.column_dimensions["B"].width = 20

    os.makedirs(outdir, exist_ok=True)
    path = os.path.join(outdir, f"datos_{d['nombre']}_{periodo.replace(' ', '_')}.xlsx")
    wb.save(path)
    return path


def main():
    ap = argparse.ArgumentParser(description="Extrae datos de seguimiento de Mailchimp por publico.")
    ap.add_argument("--publicos", nargs="*", default=PUBLICOS)
    ap.add_argument("--desde", default="2026-03-01")
    ap.add_argument("--hasta", default="2026-04-30")
    ap.add_argument("--api-key", default=os.environ.get("MAILCHIMP_API_KEY"))
    ap.add_argument("--outdir", default="informes")
    args = ap.parse_args()

    if not args.api_key:
        sys.exit("ERROR: falta MAILCHIMP_API_KEY.")

    mc = Mailchimp(args.api_key)
    periodo = f"{args.desde} a {args.hasta}"
    print(f"Periodo: {periodo}")
    for nombre in args.publicos:
        d = procesar_publico(mc, nombre, args.desde, args.hasta)
        if d and d.get("campanas"):
            path = escribir_excel(d, periodo, args.outdir)
            print(f"  [{nombre}] OK -> {path}")
        elif d:
            print(f"  [{nombre}] sin campanas en el periodo.")
    print("Listo.")


if __name__ == "__main__":
    main()
